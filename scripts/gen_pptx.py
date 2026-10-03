#!/usr/bin/env python3
"""生成授课 PPT v2(图文并茂版)。

内容 DSL(scripts/ppt_content/ch{N}.txt):
  #META no|title|subtitle
  #TOC 条目行(重绘为双栏卡片目录)
  #SECTION 标题(带讲次进度指示)
  #QUIZ 标题 | 问题行
  #CODE 标题|src/path.py:a-b
  #PAGE 标题 | 要点行 |>>备注(重绘为卡片页)
  #ARCH 标题|层1:组件,组件;层2:… |流标|tag|tag
  #FLOW 标题|步骤1->步骤2->… |loop=N|note=…|tag
  #STATE 标题|节点,…;A->B:标签;… |tag
  #SEQ 标题 |角色1,角色2,角色3 |0->1:消息;1->2:消息|tag
  #CHART bar|图表标题|k=v,k=v |insight=…|tag
  #METRICS 标题|值^后缀^标签;… |tag
  #TERM 标题|行1;行2;… |tag
  #PYRAMID 标题|名:说明;… |tag
  #TIMELINE 标题|头:述;… |tag
  #GRID 标题|头:述;… |cols=N|tag
  #TABLE 标题|表头,列…;行1,… |hl=N|tag

输出:03-课件PPT/第N章-*.pptx(完整版)+ 03-课件PPT/讲次版/(按讲切分)
用法:uv run --with python-pptx --with matplotlib python scripts/gen_pptx.py [章号…]
"""

from __future__ import annotations

import math
import re
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).parent))
import pptx_visuals as V

W, H = V.W, V.H
OUT = Path(__file__).resolve().parents[3] / "03-课件PPT"
TMP = Path(tempfile.gettempdir())

KINDS = (
    "#TOC",
    "#SECTION",
    "#PAGE",
    "#QUIZ",
    "#CODE",
    "#ARCH",
    "#FLOW",
    "#STATE",
    "#SEQ",
    "#CHART",
    "#METRICS",
    "#TERM",
    "#PYRAMID",
    "#TIMELINE",
    "#GRID",
    "#TABLE",
)


def parse_txt(path: Path):
    meta: dict = {}
    pages: list[tuple] = []
    kind, title, lines, note, extra = "", "", [], "", ""

    def flush():
        nonlocal kind, title, lines, note, extra
        if kind:
            pages.append((kind, title, "\n".join(lines), note, extra))
        kind, title, lines, note, extra = "", "", [], "", ""

    for raw in path.read_text().splitlines():
        line = raw.rstrip()
        if line.startswith("#META"):
            _, rest = line.split(" ", 1)
            no, t, s = rest.split("|")
            meta = {"no": no, "title": t, "subtitle": s}
        elif line.startswith(KINDS):
            flush()
            m = re.match(r"#(\w+)\s*(.*)", line)
            kind = m.group(1)
            parts = m.group(2).split("|")
            title = parts[0].strip() if parts[0].strip() else ""
            extra = "|".join(parts[1:]).strip() if len(parts) > 1 else ""
        elif line.startswith(">>"):
            note = line[2:].strip()
        elif line.strip():
            lines.append(line)
    flush()
    return meta, pages


def cover(prs, no, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    V.box(s, 0, 0, W, H, fill=V.INK, radius=False)
    for r, c in (
        (520, V.PANEL),
        (380, RGBColor(0x1B, 0x37, 0x63)),
        (250, RGBColor(0x24, 0x46, 0x7A)),
    ):
        cir = s.shapes.add_shape(
            V.SHP.OVAL,
            W - Emu(int(r * 7144)),
            Emu(int(r * 2858)),
            Emu(int(r * 9525)),
            Emu(int(r * 9525)),
        )
        cir.fill.solid()
        cir.fill.fore_color.rgb = c
        cir.line.fill.background()
        cir.shadow.inherit = False
    V.box(s, 0, H - Emu(240000), W, Emu(240000), fill=V.AMBER, radius=False)
    V.box(s, V._px(90), V._px(96), V._px(58), V._px(58), fill=V.AMBER)
    V.label(s, V._px(90), V._px(90), V._px(58), V._px(58), no, 22, V.INK, True)
    V.label(
        s,
        V._px(170),
        V._px(120),
        W - V._px(700),
        V._px(70),
        "AI AGENT 全栈工程师训练营",
        15,
        V.FAINT,
        align=PP_ALIGN.LEFT,
    )
    V.label(
        s,
        V._px(170),
        V._px(210),
        W - V._px(560),
        V._px(210),
        title,
        40,
        V.WHITE,
        True,
        align=PP_ALIGN.LEFT,
        spacing=1.05,
    )
    V.label(
        s,
        V._px(170),
        V._px(440),
        W - V._px(600),
        V._px(80),
        subtitle,
        16,
        V.SILVER,
        align=PP_ALIGN.LEFT,
    )
    V.label(
        s,
        V._px(170),
        H - V._px(150),
        W - V._px(400),
        V._px(50),
        "Agent Engineering Platform · v1.3 · 审计 9.4/10 · 125 测试",
        12,
        V.FAINT,
        align=PP_ALIGN.LEFT,
    )


def toc_slide(prs, items):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    V.box(s, 0, 0, W, H, fill=V.INK, radius=False)
    V.box(s, 0, 0, V._px(10), H, fill=V.AMBER, radius=False)
    V.label(
        s,
        V._px(70),
        V._px(64),
        W - V._px(200),
        V._px(60),
        "本章目录",
        28,
        V.WHITE,
        True,
        align=PP_ALIGN.LEFT,
    )
    V.label(
        s,
        V._px(70),
        V._px(130),
        W - V._px(300),
        V._px(40),
        f"{len(items)} 个知识模块 · 含实测数据图表与代码走查",
        12,
        V.FAINT,
        align=PP_ALIGN.LEFT,
    )
    half = (len(items) + 1) // 2
    for col, chunk in enumerate((items[:half], items[half:])):
        x = V._px(90) + col * int(W * 0.48)
        for i, it in enumerate(chunk):
            y = V._px(200) + i * V._px(76)
            idx = i + 1 + col * half
            c = V.ACCENTS[idx % len(V.ACCENTS)]
            V.box(s, x, y, int(W * 0.44), V._px(62), fill=V.PANEL)
            V.box(s, x, y, V._px(5), V._px(62), fill=c, radius=False)
            V.label(
                s,
                x + V._px(18),
                y,
                V._px(44),
                V._px(62),
                f"{idx:02d}",
                15,
                c,
                True,
                align=PP_ALIGN.LEFT,
            )
            V.label(
                s,
                x + V._px(68),
                y,
                int(W * 0.44) - V._px(84),
                V._px(62),
                it,
                13,
                V.SILVER,
                align=PP_ALIGN.LEFT,
            )


def section_slide(prs, title, idx, total_sections):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    V.box(s, 0, 0, W, H, fill=V.INK, radius=False)
    for i in range(max(1, total_sections)):
        c = V.AMBER if i == idx else RGBColor(0x27, 0x40, 0x66)
        V.box(s, V._px(90) + i * V._px(46), V._px(100), V._px(34), V._px(8), fill=c, radius=False)
    V.box(s, V._px(90), V._px(250), V._px(120), V._px(6), fill=V.AMBER, radius=False)
    V.label(
        s,
        V._px(90),
        V._px(300),
        W - V._px(240),
        V._px(130),
        title,
        32,
        V.WHITE,
        True,
        align=PP_ALIGN.LEFT,
        spacing=1.05,
    )
    V.label(
        s,
        V._px(90),
        V._px(450),
        W - V._px(300),
        V._px(50),
        f"LECTURE {idx + 1:02d} / {total_sections:02d}",
        14,
        V.AMBER,
        align=PP_ALIGN.LEFT,
    )


def quiz_slide(prs, title, body):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar_h = V._px(74)
    V.box(s, 0, 0, W, bar_h, fill=V.AMBER, radius=False)
    V.label(s, V._px(32), 0, W - V._px(80), bar_h, title, 20, V.INK, True, align=PP_ALIGN.LEFT)
    qs = [l for l in body.split("\n") if l.strip()]
    for i, q in enumerate(qs[:6]):
        y = V._px(120) + i * V._px(80)
        V.box(s, V._px(120), y, W - V._px(240), V._px(62), fill=V.LIGHT)
        V.box(s, V._px(120), y, V._px(5), V._px(68), fill=V.BLUE, radius=False)
        V.label(
            s,
            V._px(146),
            y,
            V._px(44),
            V._px(68),
            f"Q{i + 1}",
            15,
            V.BLUE,
            True,
            align=PP_ALIGN.LEFT,
        )
        text = re.sub(r"^Q\d+[.:、)]?\s*", "", q).strip()
        V.label(
            s, V._px(200), y, W - V._px(380), V._px(68), text, 13.5, V.GREY, align=PP_ALIGN.LEFT
        )


def card_page(prs, title, body, note, ch_no):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    V.header(s, title, f"第{ch_no}章")
    items = []
    for l in body.split("\n"):
        if not l.strip():
            continue
        t = re.sub(r"^[•▎·]+\s*", "", l).strip()
        t = re.sub(r"^\d+[.、)]\s*", "", t).strip()
        items.append(t)
    margin, gap, top = V._px(92), V._px(18), V._px(96)
    avail = H - top - V._px(104 if note else 36)
    two_col = len(items) >= 5
    if two_col:
        rows = math.ceil(len(items) / 2)
        cw = (W - margin * 2 - gap) // 2
        ch = min(V._px(92), (avail - gap * (rows - 1)) // rows)
        for i, text in enumerate(items):
            r, c = divmod(i, 2)
            x = margin + c * (cw + gap)
            y = top + r * (ch + gap)
            color = V.ACCENTS[i % len(V.ACCENTS)]
            V.box(s, x, y, cw, ch, fill=V.WHITE, line=V.BORDER)
            V.box(s, x, y, V._px(5), ch, fill=color, radius=False)
            V.label(
                s,
                x + V._px(46),
                y,
                cw - V._px(64),
                ch,
                text,
                13,
                V.GREY,
                True,
                align=PP_ALIGN.LEFT,
                spacing=1.12,
            )
            V.box(s, x + V._px(14), y + ch // 2 - V._px(7), V._px(14), V._px(14), fill=color)
    else:
        rows = max(1, len(items))
        ch = min(V._px(108), (avail - gap * (rows - 1)) // rows)
        for i, text in enumerate(items):
            y = top + i * (ch + gap)
            color = V.ACCENTS[i % len(V.ACCENTS)]
            V.box(s, margin, y, W - margin * 2, ch, fill=V.WHITE, line=V.BORDER)
            V.box(s, margin, y, V._px(6), ch, fill=color, radius=False)
            V.box(s, margin + V._px(18), y + ch // 2 - V._px(15), V._px(30), V._px(30), fill=color)
            V.label(
                s,
                margin + V._px(18),
                y + ch // 2 - V._px(16),
                V._px(30),
                V._px(30),
                str(i + 1),
                13,
                V.WHITE,
                True,
            )
            V.label(
                s,
                margin + V._px(64),
                y,
                W - margin * 2 - V._px(88),
                ch,
                text,
                14,
                V.GREY,
                True,
                align=PP_ALIGN.LEFT,
                spacing=1.12,
            )
    if note:
        y = H - V._px(88)
        V.box(s, margin, y, W - margin * 2, V._px(54), fill=V.LIGHT)
        V.label(
            s,
            margin + V._px(20),
            y,
            W - margin * 2 - V._px(40),
            V._px(58),
            "▶ " + note,
            12.5,
            V.BLUE,
            align=PP_ALIGN.LEFT,
        )


def _grab(extra, key, cast=str):
    m = re.search(rf"\|{key}=([^|]*)", extra or "")
    if not m:
        return None, extra or ""
    val = m.group(1)
    return (cast(val) if cast is not str else val), extra.replace(m.group(0), "")


def _strip(extra: str) -> str:
    return (extra or "").strip(" |")


def render_page(prs, page, ch_no, idx, total, sections_total=None):
    kind, title, body, note, extra = page
    data = body if body.strip() else extra
    if kind == "TOC":
        toc_slide(prs, [l for l in body.split("\n") if l.strip()])
    elif kind == "SECTION":
        return "SECTION", title
    elif kind == "QUIZ":
        quiz_slide(prs, title, body)
    elif kind == "PAGE":
        card_page(prs, title, body, note, ch_no)
    elif kind == "ARCH":
        layers_raw, flows, tag = data, None, ""
        m = re.search(r"\|flow=([^|]*)", layers_raw)
        if m:
            flows = m.group(1)
            layers_raw = layers_raw.replace(m.group(0), "")
        tag = layers_raw.split("|")[-1].strip() if "|" in layers_raw else ""
        if tag and not re.search(r"[\u4e00-\u9fffA-Za-z]", tag.split(":")[-1]):
            layers_raw = layers_raw.rsplit("|", 1)[0]
        layers = [
            [c.strip() for c in g.split(",")]
            for g in layers_raw.replace("\n", "").split(";")
            if g.strip()
        ]
        V.arch_diagram(
            prs.slides.add_slide(prs.slide_layouts[6]), title, layers, _strip(tag), flows or ""
        )
    elif kind == "FLOW":
        src = data
        loop, fnote = None, ""
        m = re.search(r"\|loop=(\d+)", src)
        if m:
            loop = int(m.group(1))
            src = src.replace(m.group(0), "")
        m = re.search(r"\|note=([^|]*)", src)
        if m:
            fnote = m.group(1)
            src = src.replace(m.group(0), "")
        tag = src.split("|")[-1].strip() if "|" in src else ""
        if tag and ("=" in tag or not tag):
            tag = ""
            src = src.rsplit("|", 1)[0] if "|" in src else src
        steps = [x.strip() for x in src.replace("\n", "").split("->") if x.strip()]
        V.flow_diagram(
            prs.slides.add_slide(prs.slide_layouts[6]), title, steps, tag, loop, fnote or note or ""
        )
    elif kind == "STATE":
        spec = (
            data.rsplit("|", 1)[0].strip()
            if data.count("|") and not data.rsplit("|", 1)[1].strip()
            else data
        )
        V.state_diagram(
            prs.slides.add_slide(prs.slide_layouts[6]), title, spec.replace("\n", ""), ""
        )
    elif kind == "SEQ":
        actors_part, _, msg_part = data.partition("|")
        actors = [a.strip() for a in actors_part.split(",") if a.strip()]
        msgs = []
        for m in msg_part.split(";"):
            m = m.strip()
            if not m:
                continue
            route, _, lab = m.partition(":")
            fi, _, ti = route.partition("->")
            msgs.append((int(fi), int(ti), lab.strip()))
        V.seq_diagram(prs.slides.add_slide(prs.slide_layouts[6]), title, actors, msgs)
    elif kind == "CHART":
        insight = ""
        m = re.search(r"\|insight=([^|]*)", data)
        if m:
            insight = m.group(1)
            data = data.replace(m.group(0), "")
        slide_title = title
        if "|" in data and not data.startswith(("bar", "line", "donut", "hbar", "groupbar")):
            pass
        base = f"{title}|{data}"
        spec = "|".join(base.split("|")[:3])
        png = V.chart_png(spec.replace("\n", ""), TMP / f"chart_{abs(hash(spec)) & 0xFFFFF}")
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        seg_title = spec.split("|")[1] if spec.count("|") >= 1 else title
        V.chart_slide(prs, slide, seg_title, png, "", insight or note or "")
    elif kind == "METRICS":
        cards = []
        for part in data.replace("\n", "").split(";"):
            part = part.strip()
            if part:
                v, s_, lab = (part.split("^") + ["", ""])[:3]
                cards.append((v, s_, lab))
        V.metric_cards(prs.slides.add_slide(prs.slide_layouts[6]), title, cards, _strip(extra))
    elif kind == "TERM":
        src = data.split("|")[0]
        lines = src.split("\n") if "\n" in src else src.split(";")
        V.terminal(
            prs.slides.add_slide(prs.slide_layouts[6]),
            title,
            [l for l in lines if l.strip()],
            _strip(extra),
        )
    elif kind == "PYRAMID":
        levels = []
        for g in data.replace("\n", "").split(";"):
            if g.strip():
                n_, _, d_ = g.partition(":")
                levels.append((n_.strip(), d_.strip()))
        V.pyramid(prs.slides.add_slide(prs.slide_layouts[6]), title, levels, _strip(extra))
    elif kind == "TIMELINE":
        items = []
        for g in data.replace("\n", "").split(";"):
            if g.strip():
                h_, _, d_ = g.partition(":")
                items.append((h_.strip(), d_.strip()))
        V.timeline(prs.slides.add_slide(prs.slide_layouts[6]), title, items, _strip(extra))
    elif kind == "GRID":
        cols = 3
        m = re.search(r"\|cols=(\d)", data)
        if m:
            cols = int(m.group(1))
            data = data.replace(m.group(0), "")
        cards = []
        for g in data.replace("\n", "").split(";"):
            if g.strip():
                h_, _, d_ = g.partition(":")
                cards.append((h_.strip(), d_.strip()))
        V.card_grid(
            prs.slides.add_slide(prs.slide_layouts[6]), title, cards, _strip(""), cols=cols or 3
        )
    elif kind == "TABLE":
        hl = -1
        m = re.search(r"\|hl=(\d+)", data)
        if m:
            hl = int(m.group(1))
            data = data.replace(m.group(0), "")
        tbl_src = data.split("|")[0]
        rows = [
            [c.strip() for c in ln.split(",")]
            for ln in tbl_src.replace("\n", "").split(";")
            if ln.strip()
        ]
        V.styled_table(
            prs.slides.add_slide(prs.slide_layouts[6]),
            title,
            rows,
            _strip(""),
            highlight_col=hl if hl is not None else -1,
        )
    elif kind == "CODE":
        spec = extra if extra.strip() else body
        path_part, _, range_part = spec.partition(":")
        f = Path(__file__).resolve().parents[1] / path_part.strip()
        a, _, b = range_part.partition("-")
        src = f.read_text().splitlines()
        code = "\n".join(src[int(a) - 1 : int(b or a)])
        png = TMP / f"code_{abs(hash(body)) & 0xFFFFF}.png"
        w, h = _render_code(code, png)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        V.box(slide, 0, 0, W, H, fill=V.INK, radius=False)
        V.box(slide, 0, 0, W, V._px(64), fill=V.PANEL, radius=False)
        V.label(
            slide,
            V._px(28),
            0,
            W - V._px(80),
            V._px(64),
            f"{title} · {path_part}:{a}-{b}",
            16,
            V.AMBER,
            True,
            align=PP_ALIGN.LEFT,
        )
        scale = min(int(W * 0.94) / w, int(H * 0.84) / h)
        slide.shapes.add_picture(
            str(png), int(W - w * scale) // 2, V._px(80), int(w * scale), int(h * scale)
        )
    last = prs.slides[len(prs.slides._sldIdLst) - 1]
    V.footer(last, idx, total, ch_no)
    return None


def _render_code(code: str, out_png: Path):
    from PIL import Image, ImageDraw, ImageFont

    lines = code.splitlines()
    img = Image.new("RGB", (1760, max(220, 90 + 30 * len(lines) + 40)), (13, 27, 42))
    d = ImageDraw.Draw(img)
    mono = ImageFont.truetype("/System/Library/Fonts/Menlo.ttc", 19)
    cjk = ImageFont.truetype("/System/Library/Fonts/Hiragino Sans GB.ttc", 19)
    gutter = ImageFont.truetype("/System/Library/Fonts/Menlo.ttc", 15)
    y = 60
    for i, ln in enumerate(lines, 1):
        ln = ln[:96]
        d.text((56, y), f"{i:>3}", font=gutter, fill=(94, 110, 130))
        font = cjk if any(ord(c) > 0x2E80 for c in ln) else mono
        color = (200, 210, 225)
        st = ln.lstrip()
        if st.startswith(("class ", "def ", "async def")):
            color = (120, 180, 255)
        elif st.startswith("#"):
            color = (120, 200, 140)
        d.text((110, y), ln, font=font, fill=color)
        y += 30
    img.save(out_png)
    return img.size


def apply_lecture_branding(prs, lec_no: int, lec_title: str):
    """教研室模板:封面代码化徽章 + 全页左下页眉标签(仅讲次版)。"""
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as _SHP
    # 封面徽章(双环 + AEP + 说明)
    cover = prs.slides[0]
    cx, cy, r = W - V._px(240), H - V._px(190), V._px(56)
    ring = cover.shapes.add_shape(_SHP.OVAL, cx - r, cy - r, r * 2, r * 2)
    ring.fill.background(); ring.line.color.rgb = V.AMBER
    ring.line.width = __import__("pptx").util.Pt(2)
    ring.shadow.inherit = False
    inner = cover.shapes.add_shape(_SHP.OVAL, cx - r + V._px(8), cy - r + V._px(8),
                                   (r - V._px(8)) * 2, (r - V._px(8)) * 2)
    inner.fill.solid(); inner.fill.fore_color.rgb = V.PANEL
    inner.line.fill.background(); inner.shadow.inherit = False
    V.label(cover, cx - r, cy - V._px(16), r * 2, V._px(32), "AEP", 18, V.AMBER, True)
    V.label(cover, cx - r - V._px(60), cy + r + V._px(6), r * 2 + V._px(120), V._px(28),
            "AI Agent 训练营 · 教研室", 10, V.FAINT)
    # 其余页:左下页眉
    for sl in list(prs.slides)[1:]:
        V.label(sl, V._px(64), H - V._px(34), V._px(560), V._px(28),
                f"教研室内部课件 · 讲次 {lec_no:02d} · {lec_title[:18]}", 9, V.FAINT,
                align=__import__("pptx").enum.text.PP_ALIGN.LEFT)


def build(module_name: str, lecture_split: bool = True):
    txt = Path(__file__).parent / "ppt_content" / f"{module_name}.txt"
    meta, pages = parse_txt(txt)
    sections = [i for i, p in enumerate(pages) if p[0] == "SECTION"]
    section_titles = [pages[i][1] for i in sections]
    total = len(pages) + 1
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    cover(prs, meta["no"], meta["title"], meta["subtitle"])
    sec_seen = 0
    section_idx_map: dict[str, int] = {}
    for i, page in enumerate(pages, start=2):
        result = render_page(prs, page, meta["no"], i, total)
        if result and result[0] == "SECTION":
            section_idx_map[result[1]] = sec_seen
            # 重绘 section 页(带讲次进度)
            sl = prs.slides[len(prs.slides._sldIdLst) - 1]
            section_slide(prs, result[1], sec_seen, len(sections))
            sec_seen += 1
    safe = meta["title"].split(":")[0].split("、")[0]
    out = OUT / f"第{meta['no']}章-{safe}.pptx"
    prs.save(out)
    print(f"完整版: {out.name}({len(prs.slides._sldIdLst)} 页)")

    if lecture_split and sections:
        bounds = list(zip(sections, sections[1:] + [len(pages)]))
        (OUT / "讲次版").mkdir(exist_ok=True)
        for li, (start, end) in enumerate(bounds):
            lp = Presentation()
            lp.slide_width, lp.slide_height = W, H
            sub = pages[start][1]
            cover(lp, f"{meta['no']}.{li + 1}", sub.replace("讲次 ", ""), meta["title"])
            seg = pages[:1] + pages[start:end]
            for j, page in enumerate(seg, start=2):
                res = render_page(lp, page, meta["no"], j, len(seg) + 1)
                if res and res[0] == "SECTION":
                    section_slide(lp, res[1], 0, 1)
            apply_lecture_branding(lp, li + 1, sub.replace("讲次 ", ""))
            name = re.sub(r"[\\/:*?\"<>|]", "", sub)[:28]
            lo = OUT / "讲次版" / f"第{meta['no']}章-{li + 1}-{name}.pptx"
            lp.save(lo)
        print(f"  讲次版: {len(bounds)} 个文件")


def main() -> int:
    targets = sys.argv[1:] or [f"ch{i}" for i in range(1, 8)]
    OUT.mkdir(parents=True, exist_ok=True)
    for t in targets:
        build(t)
    return 0


if __name__ == "__main__":
    sys.exit(main())
